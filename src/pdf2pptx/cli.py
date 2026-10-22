import argparse
from pathlib import Path

try:
    import pymupdf as fitz
except ImportError:
    import fitz

from pptx import Presentation
from pptx.util import Emu

from pdf2pptx import __version__

SLIDE_WIDTH_EMU = 12192000


def _quality(value):
    q = int(value)
    if not 1 <= q <= 100:
        raise argparse.ArgumentTypeError("quality must be between 1 and 100")
    return q


def build_parser():
    parser = argparse.ArgumentParser(
        prog="pdf2pptx",
        description="Render every PDF page to a high-quality image and build a PowerPoint with one full-bleed slide per page.",
    )
    parser.add_argument("pdf", help="input PDF file")
    parser.add_argument("-o", "--output", help="output .pptx (default: <pdf name>.pptx)")
    parser.add_argument("--dpi", type=int, default=300, help="image resolution in DPI relative to slide width (default: 300)")
    parser.add_argument("--format", choices=["jpeg", "png"], default="jpeg", help="image format (default: jpeg)")
    parser.add_argument("--quality", type=_quality, default=95, help="JPEG quality, 1-100 (default: 95)")
    parser.add_argument("--image-dir", help="folder for rendered pages (default: <pdf name>_pages)")
    parser.add_argument("--no-keep-images", dest="keep_images", action="store_false", help="delete rendered images after building the pptx")
    parser.set_defaults(keep_images=True)
    parser.add_argument("--version", action="version", version=f"%(prog)s {__version__}")
    return parser


def convert(pdf_path, output=None, dpi=300, image_format="jpeg", quality=95, image_dir=None, keep_images=True):
    pdf_path = Path(pdf_path).expanduser().resolve()
    if not pdf_path.is_file():
        raise FileNotFoundError(f"PDF not found: {pdf_path}")

    out_path = Path(output).expanduser().resolve() if output else pdf_path.with_suffix(".pptx")
    if out_path.suffix.lower() != ".pptx":
        out_path = out_path.with_suffix(".pptx")

    image_dir = Path(image_dir).expanduser().resolve() if image_dir else pdf_path.with_name(f"{pdf_path.stem}_pages")
    image_dir.mkdir(parents=True, exist_ok=True)

    doc = fitz.open(pdf_path)
    if doc.page_count == 0:
        raise ValueError("PDF has no pages")

    page_rect = doc[0].rect
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_WIDTH_EMU)
    prs.slide_height = Emu(round(SLIDE_WIDTH_EMU * page_rect.height / page_rect.width))
    slide_width_in = prs.slide_width / 914400
    slide_height_in = prs.slide_height / 914400
    blank_layout = prs.slide_layouts[6]
    ext = "jpg" if image_format == "jpeg" else "png"

    print(f"PDF: {pdf_path} ({doc.page_count} pages)")
    images = []
    try:
        for i, page in enumerate(doc, start=1):
            zoom_x = dpi * slide_width_in / page.rect.width
            zoom_y = dpi * slide_height_in / page.rect.height
            pix = page.get_pixmap(matrix=fitz.Matrix(zoom_x, zoom_y), colorspace=fitz.csRGB, alpha=False)

            img_path = image_dir / f"page-{i:04d}.{ext}"
            if image_format == "jpeg":
                img_path.write_bytes(pix.tobytes("jpeg", jpg_quality=quality))
            else:
                img_path.write_bytes(pix.tobytes("png"))
            images.append(img_path)
            print(f"  page {i}/{doc.page_count}: {pix.width}x{pix.height}px -> {img_path}")

            slide = prs.slides.add_slide(blank_layout)
            slide.shapes.add_picture(str(img_path), 0, 0, prs.slide_width, prs.slide_height)

        prs.save(out_path)
        print(f"Saved: {out_path} ({doc.page_count} slides)")
    finally:
        doc.close()

    if not keep_images:
        for path in images:
            path.unlink(missing_ok=True)
        try:
            image_dir.rmdir()
        except OSError:
            pass
        print(f"Removed rendered images from {image_dir}")

    return out_path


def main(argv=None):
    args = build_parser().parse_args(argv)
    try:
        convert(
            args.pdf,
            output=args.output,
            dpi=args.dpi,
            image_format=args.format,
            quality=args.quality,
            image_dir=args.image_dir,
            keep_images=args.keep_images,
        )
    except (FileNotFoundError, ValueError) as exc:
        raise SystemExit(f"error: {exc}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
